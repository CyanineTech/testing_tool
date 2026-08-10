#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse
import json
import logging
import os
import sys
import time
from collections import defaultdict
from configparser import ConfigParser
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

import requests
from openpyxl import load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_CONFIG_PATH = os.path.join(BASE_DIR, "config.ini")
ENDPOINT_PATH = "/dispatch_server/dispatch/start/location_call/task/"


@dataclass
class RuleRuntime:
    name: str
    hz_per_hour: float
    interval_seconds: float
    tasks: List[Dict[str, str]]
    next_run: float
    cursor: int = 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="根据 config.ini 的 logic_mapping 规则，按频率自动调用开始任务接口"
    )
    parser.add_argument("--config", default=DEFAULT_CONFIG_PATH, help="配置文件路径")
    parser.add_argument(
        "--method",
        default="PUT",
        choices=["PUT", "POST"],
        help="调用方法，默认 PUT",
    )
    parser.add_argument(
        "--rules",
        nargs="*",
        default=None,
        help="仅运行指定规则名，例如：--rules rule_1 rule_3",
    )
    parser.add_argument(
        "--max-hours",
        type=float,
        default=None,
        help="运行满指定小时后自动停止，例如：--max-hours 2.5",
    )
    parser.add_argument("--once", action="store_true", help="每条规则只执行 1 次后退出")
    parser.add_argument("--dry-run", action="store_true", help="仅打印将要调用的任务，不发送请求")
    return parser.parse_args()


def setup_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
    )


def load_config(config_path: str) -> ConfigParser:
    if not os.path.exists(config_path):
        raise FileNotFoundError(f"配置文件不存在: {config_path}")

    cfg = ConfigParser(interpolation=None, delimiters=("=",))
    cfg.optionxform = str
    cfg.read(config_path, encoding="utf-8")
    return cfg


def get_required(cfg: ConfigParser, section: str, key: str) -> str:
    if not cfg.has_section(section):
        raise ValueError(f"缺少配置段: [{section}]")
    if not cfg.has_option(section, key):
        raise ValueError(f"缺少配置项: [{section}] {key}")
    value = cfg.get(section, key).strip()
    if not value:
        raise ValueError(f"配置为空: [{section}] {key}")
    return value


def to_abs_path(path_value: str, config_path: str) -> str:
    if os.path.isabs(path_value):
        return path_value
    config_dir = os.path.dirname(os.path.abspath(config_path))
    return os.path.join(config_dir, path_value)


def parse_rule_json(raw: str, rule_name: str) -> Tuple[List[str], List[str], float]:
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{rule_name} 不是合法JSON: {exc}") from exc

    if not isinstance(data, dict):
        raise ValueError(f"{rule_name} 必须是对象JSON")

    pick = data.get("pick")
    drop = data.get("drop")
    hz = data.get("hz", 5)

    if not isinstance(pick, list) or not pick:
        raise ValueError(f"{rule_name} 的 pick 必须是非空列表")
    if not isinstance(drop, list) or not drop:
        raise ValueError(f"{rule_name} 的 drop 必须是非空列表")

    pick_list = [str(x).strip() for x in pick if str(x).strip()]
    drop_list = [str(x).strip() for x in drop if str(x).strip()]
    if not pick_list or not drop_list:
        raise ValueError(f"{rule_name} 的 pick/drop 不能是空字符串列表")

    if "hz" not in data:
        logging.warning("%s 未配置 hz，使用默认值 5 次/小时", rule_name)

    try:
        hz_value = float(hz)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{rule_name} 的 hz 必须是数字") from exc

    if hz_value <= 0:
        raise ValueError(f"{rule_name} 的 hz 必须大于0（单位：次/小时）")

    return pick_list, drop_list, hz_value


def extract_area_from_alias(alias_kept: str) -> str:
    normalized = str(alias_kept).strip().replace(" - ", "-")
    if not normalized:
        return ""
    parts = normalized.rsplit("-", 1)
    if len(parts) == 2:
        return parts[0].strip()
    return normalized


def load_area_locations(xlsx_path: str, sheet_name: Optional[str]) -> Dict[str, List[str]]:
    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(f"Excel 文件不存在: {xlsx_path}")

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            wb.close()
            raise ValueError(f"Excel 不存在工作表: {sheet_name}")
        ws = wb[sheet_name]
    else:
        ws = wb[wb.sheetnames[-1]]

    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        wb.close()
        raise ValueError("Excel 无表头")

    headers = [str(x).strip().lower() if x is not None else "" for x in header]
    if "id" not in headers or "alias_kept" not in headers:
        wb.close()
        raise ValueError("Excel 必须包含 id 和 alias_kept 列")

    id_idx = headers.index("id")
    alias_idx = headers.index("alias_kept")

    area_map: Dict[str, List[str]] = defaultdict(list)
    seen_ids = set()

    for row in rows:
        if row is None:
            continue
        location_id = str(row[id_idx]).strip() if row[id_idx] is not None else ""
        alias_value = str(row[alias_idx]).strip() if row[alias_idx] is not None else ""

        if not location_id or not alias_value:
            continue
        if location_id in seen_ids:
            continue

        area = extract_area_from_alias(alias_value)
        if not area:
            continue

        seen_ids.add(location_id)
        area_map[area].append(location_id)

    wb.close()
    return dict(area_map)


def expand_tasks_for_rule(
    rule_name: str,
    pick_list: List[str],
    drop_list: List[str],
    area_locations: Dict[str, List[str]],
) -> List[Dict[str, str]]:
    mode = "single"
    if len(drop_list) == 1:
        mode = "single"
    elif len(drop_list) == len(pick_list):
        mode = "one_to_one"
    else:
        mode = "round_robin"
        logging.warning(
            "规则 %s: pick/drop 长度不匹配 (pick=%s, drop=%s)，将按 drop 列表轮询分配。",
            rule_name,
            len(pick_list),
            len(drop_list),
        )

    tasks: List[Dict[str, str]] = []
    round_robin_index = 0
    for pick_index, pick in enumerate(pick_list):
        location_ids = area_locations.get(pick, [])
        if not location_ids:
            logging.warning("规则 %s: pick 区域 %s 未在 Excel 找到任何 location_id", rule_name, pick)
            continue
        for location_id in location_ids:
            if mode == "single":
                target_drop = drop_list[0]
            elif mode == "one_to_one":
                target_drop = drop_list[pick_index]
            else:
                target_drop = drop_list[round_robin_index % len(drop_list)]
                round_robin_index += 1
            tasks.append({"location_id": location_id, "area": target_drop})

    return tasks


def load_rules(cfg: ConfigParser, selected_rules: Optional[List[str]]) -> List[Tuple[str, List[str], List[str], float]]:
    if not cfg.has_section("logic_mapping"):
        raise ValueError("缺少配置段: [logic_mapping]")

    rules: List[Tuple[str, List[str], List[str], float]] = []
    for key, value in cfg.items("logic_mapping"):
        if not key.startswith("rule_"):
            continue
        if selected_rules and key not in selected_rules:
            continue
        pick_list, drop_list, hz_value = parse_rule_json(value, key)
        rules.append((key, pick_list, drop_list, hz_value))

    if not rules:
        raise ValueError("未找到可执行规则，请检查 [logic_mapping] 或 --rules 参数")

    rules.sort(key=lambda item: item[0])
    return rules


def send_task(
    session: requests.Session,
    method: str,
    url: str,
    payload: Dict[str, str],
    headers: Dict[str, str],
    timeout: float,
    retry_count: int,
    retry_delay: float,
    dry_run: bool,
) -> bool:
    if dry_run:
        logging.info("[DRY RUN] %s %s payload=%s", method, url, json.dumps(payload, ensure_ascii=False))
        return True

    max_attempts = retry_count + 1
    for attempt in range(1, max_attempts + 1):
        try:
            resp = session.request(method=method, url=url, json=payload, headers=headers, timeout=timeout)
            ok = 200 <= resp.status_code < 300
            if ok:
                logging.info("请求成功 status=%s payload=%s", resp.status_code, json.dumps(payload, ensure_ascii=False))
                return True

            logging.error("请求失败 status=%s body=%s", resp.status_code, resp.text[:300])
        except requests.RequestException as exc:
            logging.error("请求异常: %s", exc)

        if attempt < max_attempts:
            time.sleep(max(retry_delay, 0.0))

    return False


def run_scheduler(
    runtimes: List[RuleRuntime],
    session: requests.Session,
    method: str,
    api_url: str,
    headers: Dict[str, str],
    timeout: float,
    retry_count: int,
    retry_delay: float,
    once: bool,
    dry_run: bool,
    max_hours: Optional[float],
) -> None:
    total_sent = 0
    total_failed = 0
    stop_reason = "unknown"
    start_time = time.monotonic()
    max_duration_seconds = max_hours * 3600.0 if max_hours is not None else None
    rule_stats: Dict[str, Dict[str, int]] = {
        runtime.name: {"sent": 0, "success": 0, "failed": 0}
        for runtime in runtimes
    }

    logging.info("启动规则调度器，共 %s 条规则", len(runtimes))
    for runtime in runtimes:
        logging.info(
            "规则 %s: hz=%s次/小时, 间隔=%ss, 任务池=%s",
            runtime.name,
            runtime.hz_per_hour,
            runtime.interval_seconds,
            len(runtime.tasks),
        )

    try:
        while True:
            now = time.monotonic()

            if max_duration_seconds is not None and (now - start_time) >= max_duration_seconds:
                stop_reason = f"达到最大运行时长 {max_hours} 小时"
                break

            triggered = 0

            for runtime in runtimes:
                if now < runtime.next_run:
                    continue

                payload = runtime.tasks[runtime.cursor % len(runtime.tasks)]
                runtime.cursor += 1
                runtime.next_run += runtime.interval_seconds

                ok = send_task(
                    session=session,
                    method=method,
                    url=api_url,
                    payload=payload,
                    headers=headers,
                    timeout=timeout,
                    retry_count=retry_count,
                    retry_delay=retry_delay,
                    dry_run=dry_run,
                )
                total_sent += 1
                rule_stats[runtime.name]["sent"] += 1
                if not ok:
                    total_failed += 1
                    rule_stats[runtime.name]["failed"] += 1
                else:
                    rule_stats[runtime.name]["success"] += 1

                triggered += 1

            if once:
                stop_reason = "--once 模式：每条规则执行 1 次后停止"
                break

            if triggered == 0:
                time.sleep(0.2)
    except KeyboardInterrupt:
        stop_reason = "收到 Ctrl+C 中断信号"
        logging.info("收到中断信号，准备退出")

    if stop_reason == "unknown":
        stop_reason = "调度循环自然结束"

    success = total_sent - total_failed
    logging.info("调度结束，总请求=%s, 成功=%s, 失败=%s", total_sent, success, total_failed)
    logging.info("停止原因：%s", stop_reason)

    logging.info("各规则调用统计：")
    for runtime in runtimes:
        stats = rule_stats[runtime.name]
        sent = stats["sent"]
        success_count = stats["success"]
        failed_count = stats["failed"]
        success_rate = (success_count / sent * 100.0) if sent > 0 else 0.0
        logging.info(
            "  %s -> 总调用=%s, 成功=%s, 失败=%s, 成功率=%.2f%%",
            runtime.name,
            sent,
            success_count,
            failed_count,
            success_rate,
        )


def main() -> int:
    args = parse_args()
    setup_logging()

    try:
        cfg = load_config(args.config)

        host = get_required(cfg, "service", "host")
        port_str = get_required(cfg, "service", "port")
        token = get_required(cfg, "base", "token")

        try:
            port = int(port_str)
        except ValueError as exc:
            raise ValueError(f"[service] port 不是整数: {port_str}") from exc

        xlsx_path_cfg = get_required(cfg, "excel", "xlsx_path")
        xlsx_path = to_abs_path(xlsx_path_cfg, args.config)
        sheet_name = cfg.get("excel", "sheet_name", fallback="").strip() or None

        timeout = float(cfg.get("request", "timeout", fallback="15.0"))
        retry_count = int(cfg.get("request", "retry_count", fallback="0"))
        retry_delay = float(cfg.get("request", "retry_delay", fallback="1.0"))

        rules = load_rules(cfg, args.rules)
        area_locations = load_area_locations(xlsx_path, sheet_name)

        runtimes: List[RuleRuntime] = []
        for rule_name, pick_list, drop_list, hz_value in rules:
            tasks = expand_tasks_for_rule(rule_name, pick_list, drop_list, area_locations)
            if not tasks:
                logging.warning("规则 %s 没有可执行任务，跳过", rule_name)
                continue
            interval_seconds = 3600.0 / hz_value
            runtimes.append(
                RuleRuntime(
                    name=rule_name,
                    hz_per_hour=hz_value,
                    interval_seconds=interval_seconds,
                    tasks=tasks,
                    next_run=time.monotonic(),
                )
            )

        if not runtimes:
            raise RuntimeError("没有可执行规则任务，请检查 logic_mapping 与 Excel 匹配")

        api_url = f"http://{host}:{port}{ENDPOINT_PATH}"
        headers: Dict[str, str] = {
            "Content-Type": "application/json",
            "Authorization": f"Bearer {token}",
        }

        logging.info("接口地址: %s", api_url)
        logging.info("Excel路径: %s", xlsx_path)

        with requests.Session() as session:
            run_scheduler(
                runtimes=runtimes,
                session=session,
                method=args.method,
                api_url=api_url,
                headers=headers,
                timeout=timeout,
                retry_count=max(retry_count, 0),
                retry_delay=retry_delay,
                once=args.once,
                dry_run=args.dry_run,
                max_hours=args.max_hours,
            )

        return 0
    except Exception as exc:
        logging.error("执行失败: %s", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
